"""Модули заявлений (0.7.3): папка forms/<id>/, хуки бланка, чтение эталона по подписям бланка (reference:),
загрузчики DOCX (порядок документа) и PDF (значения, вписанные поверх бланка)."""
import shutil
import sys
from pathlib import Path

import pytest
import yaml

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from doctool import application, formspec, reference  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
FORM = formspec.form_by_id("regru_transfer_person")

TEMPLATE = """Заявление в ООО «Регистратор доменных имен РЕГ.РУ» от Администратора домена - физического лица о передаче
права администрирования домена другому лицу
Генеральному директору
ООО «Регистратор доменных имен РЕГ.РУ»
от ___Тестова Петра Сергеевича, 15.03.1990___
(ФИО, дата рождения)
__________________________________
(ИНН ИП (при наличии статуса ИП))
4509 123456, 20.04.2010
(серия, номер паспорта, когда выдан)
ГУ МВД России по г. Москве
{кем выдан)
зарегистрированного по адресу
г. Москва, ул. Тестовая, д. 1
(адрес регистрации)
Заявление
Я, Тестов Пётр Сергеевич
(фамилия, имя, отчество)
прошу передать права по администрированию домена(ов)
alpha-test.ru, business-
dvd.ru
(наименование домена(ов)
и предоставляемых по этим доменам дополнительных услуг (при необходимости)
________________________
(Primary, Secondary DNS)
новому Администратору:
Для юр.лиц: | Наименование юр.лица | ООО «Тест»
e-mail/тел. | info@test.ru
Для физ.лиц и ИП: | ФИО |
e-mail/тел. |
Номер договора (название аккаунта) нового администратора / партнера (с указанием профиля): 1234567
(номер договора с компанией РЕГ.РУ)
____________ / Тестов П. С. /
(подпись) (ФИО)
Первое октября две тысячи двадцать шестого года
(дата прописью)
"""


def test_module_layout_and_hooks():
    d = formspec.module_dir(FORM)
    assert d == ROOT / "forms" / "regru_transfer_person" and (d / "form.yaml").exists()
    h = formspec.hooks(FORM)
    assert callable(h.after_text) and callable(h.scan_page) and callable(h.parse_new_admin_block)
    assert formspec.call_hook(None, "after_text", default="нет") == "нет"
    with pytest.raises(KeyError):
        formspec.call_hook(FORM, "нет_такого_хука")
    # раздел «новому Администратору» — в модуле бланка, не в ядре
    assert not hasattr(application, "apply_new_admin_block") and not hasattr(application, "parse_new_admin_block")


def test_literals():
    rx = reference.literal_rx
    assert rx("(кем выдан)").search("… {кем  выдан) …")                 # вид скобок и пробелы не важны
    assert rx("от").search("ООО «Регистратор» от Иванова") and not rx("от").search("работы")   # целое слово
    assert rx("e-mail/тел.").search("E-MAIL/ТЕЛ.")


def test_reference_reading():
    ex = application.extract_from_text(TEMPLATE, FORM, "x.pdf")
    assert ex.debug["reader"] == "reference" and ex.debug["reference"]["found"] == ex.debug["reference"]["total"]
    assert ex.get("applicant_header.raw_fio") == "Тестова Петра Сергеевича"
    assert ex.get("applicant_header.birth_date") == "15.03.1990" and not ex.get("inn")
    assert ex.get("passport") == "4509 123456" and ex.get("passport.issue_date") == "20.04.2010"
    assert ex.get("issued_by") == "ГУ МВД России по г. Москве" and ex.get("address") == "г. Москва, ул. Тестовая, д. 1"
    assert ex.get("domains.list") == ["alpha-test.ru", "business-dvd.ru"]          # перенос имени после дефиса
    assert ex.get("new_admin_org") == "ООО «Тест»" and ex.get("new_admin_org_contact.emails") == ["info@test.ru"]
    assert not ex.get("new_admin_fio") and not ex.get("services")
    assert ex.get("contract") == "1234567" and ex.get("application_date") == "01.10.2026"
    # OCR-текст скана по подписям не читается — только регулярки
    assert application.extract_from_text(TEMPLATE, FORM, "x.png", conf=0.85, source="ocr-text").debug.get("reader") is None
    # не бланк (подписей меньше min_found) — регулярки, как раньше
    free = "Генеральному директору\nот Тестова Петра Сергеевича, 15.03.1990\nПаспорт 4509 123456\n"
    ex2 = application.extract_from_text(free, FORM, "y.pdf")
    assert ex2.debug.get("reader") is None and ex2.get("passport") == "4509 123456"


def test_reference_value_that_does_not_parse_falls_back_to_regex():
    # подпись на месте, но в строке паспорта не паспорт — поле читают регулярки (здесь — из строки «Паспорт …» ниже)
    text = TEMPLATE.replace("4509 123456, 20.04.2010\n(серия", "см. ниже\n(серия") + "Паспорт 4510 654321\n"
    ex = application.extract_from_text(text, FORM, "x.pdf")
    assert ex.debug["reader"] == "reference" and ex.get("passport") == "4510 654321"


def test_docx_document_order_and_merged_cells(tmp_path):
    import docx
    from doctool.ocr import docx_text
    d = docx.Document()
    d.add_paragraph("новому Администратору:")
    t = d.add_table(rows=2, cols=3)
    for i, row in enumerate([("Для физ.лиц и ИП:", "ФИО", "Тестов Пётр"), ("", "e-mail/тел.", "t@example.com")]):
        for j, v in enumerate(row):
            t.cell(i, j).text = v
    t.cell(0, 0).merge(t.cell(1, 0)).text = "Для физ.лиц и ИП:"
    d.add_paragraph("Номер договора: 1")
    d.save(tmp_path / "a.docx")
    assert docx_text(tmp_path / "a.docx").splitlines() == [
        "новому Администратору:", "Для физ.лиц и ИП: | ФИО | Тестов Пётр", "e-mail/тел. | t@example.com", "Номер договора: 1"]


def test_pdf_values_written_over_blank(tmp_path):
    pymupdf = pytest.importorskip("pymupdf")
    from doctool.ocr import pdf_page_text
    font = next((p for p in ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "C:/Windows/Fonts/arial.ttf")
                 if Path(p).exists()), None)
    if font is None:
        pytest.skip("нет шрифта с кириллицей")
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    page.insert_font(fontname="f", fontfile=font)
    for i, line in enumerate(["от ______________________", "(ФИО, дата рождения)", "______________________", "(кем выдан)"]):
        page.insert_text((60, 100 + 20 * i), line, fontname="f", fontsize=10)
    page.add_freetext_annot(pymupdf.Rect(80, 89, 400, 104), "Тестова Петра Сергеевича", fontsize=9)
    w = pymupdf.Widget()
    w.field_type, w.field_name, w.field_value = pymupdf.PDF_WIDGET_TYPE_TEXT, "issued_by", "ГУ МВД России"
    w.rect = pymupdf.Rect(60, 132, 400, 146)
    page.add_widget(w)
    doc.save(tmp_path / "o.pdf")
    lines = pdf_page_text(pymupdf.open(tmp_path / "o.pdf")[0]).splitlines()
    i = next(n for n, x in enumerate(lines) if "Тестова" in x)
    assert lines[i + 1] == "(ФИО, дата рождения)" and "ГУ МВД России" in lines[i + 2]   # значения — у своих подписей


def test_module_folder_wins_over_stale_file(tmp_path, monkeypatch):
    forms = tmp_path / "forms"
    shutil.copytree(ROOT / "forms" / "regru_transfer_person", forms / "regru_transfer_person")
    data = yaml.safe_load((forms / "regru_transfer_person" / "form.yaml").read_text(encoding="utf-8"))
    data["title"] = "устаревшая копия"
    (forms / "regru_transfer_person.yaml").write_text(yaml.safe_dump(data, allow_unicode=True), encoding="utf-8")
    loaded = formspec.load_forms(forms)
    assert len(loaded) == 1 and loaded[0]["title"] != "устаревшая копия" and loaded[0]["_module"]
    assert formspec.stale_files(forms) == ["regru_transfer_person.yaml"]
    assert formspec.validate_config(forms) == []


def test_validate_module(tmp_path):
    forms = tmp_path / "forms"
    mod = forms / "regru_transfer_person"
    shutil.copytree(ROOT / "forms" / "regru_transfer_person", mod)
    data = yaml.safe_load((mod / "form.yaml").read_text(encoding="utf-8"))
    data["reference"]["fields"]["нет_поля"] = {"before": "(x)"}
    data["reference"]["fields"]["inn"] = {"after": 5}
    data["reference"]["min_found"] = 2
    (mod / "form.yaml").write_text(yaml.safe_dump(data, allow_unicode=True), encoding="utf-8")
    (mod / "hooks.py").write_text("import нет_такого_модуля\n", encoding="utf-8")
    (mod / "samples" / "lost.expected.json").write_text('{"fields": {"неизвестное": 1}}', encoding="utf-8")
    problems = formspec.validate_config(forms)
    text = "\n".join(problems)
    for part in ("поля 'нет_поля' нет", "поле inn: нужна подпись before", "min_found", "hooks.py не загружается",
                 "lost.expected.json: нет файла образца", "поле 'неизвестное' не найдено"):
        assert part in text, (part, problems)


def test_blank_docx_checks_reference_literals(tmp_path):
    import docx
    forms = tmp_path / "forms"
    mod = forms / "regru_transfer_person"
    shutil.copytree(ROOT / "forms" / "regru_transfer_person", mod)
    d = docx.Document()
    for line in TEMPLATE.replace("(подпись) (ФИО)", "(подпись)").splitlines():   # в бланке нет «(ФИО)» у подписи
        d.add_paragraph(line)
    d.save(mod / "blank.docx")
    problems = formspec.validate_config(forms)
    assert len(problems) == 1 and "blank.docx" in problems[0] and "«(ФИО)»" in problems[0]


def test_export_reader():
    from types import SimpleNamespace

    from doctool.service import _fields_reader
    assert _fields_reader(SimpleNamespace(app=application.extract_from_text(TEMPLATE, FORM, "x.pdf"))) == "reference"
    free = application.extract_from_text("Генеральному директору\nПаспорт 4509 123456\n", FORM, "y.pdf")
    assert _fields_reader(SimpleNamespace(app=free)) == "regex"
